#!/usr/bin/env python3
"""
Update the Attendance sheet in grades.xlsx from Moodle retrieval-quiz
grade report(s).

Marks each roster student 'p' (present — completed the quiz) or 'n'
(not present — no completed attempt) in the date column matching the
quiz's own timestamps. The target date is read directly from each CSV's
"Started" column, not guessed from the filename, so it works for any
week/day without hardcoding a Mon/Wed/Fri schedule.

Cells that already hold a value are left untouched — a manually-entered
'p' (e.g. a justified absence) or anything else already in the sheet
always wins over what the quiz report says. Only blank cells are filled
in. Use --force if a cell genuinely needs to be recomputed.

Usage:
    cd grades/

    # Process every "26FA DATA-202...-grades.csv" report found in this folder:
    python3 update_attendance_from_quiz.py

    # Or target specific file(s) explicitly:
    python3 update_attendance_from_quiz.py "26FA DATA-202-A-Retrieval Quiz W01D2-grades.csv"

    # Preview without writing anything:
    python3 update_attendance_from_quiz.py --dry-run
"""

import csv
import glob
import os
import shutil
import sys
from collections import Counter
from datetime import datetime

import openpyxl

ATTENDANCE_SHEET = "Attendance"
QUIZ_CSV_GLOB = "26FA DATA-202*.csv"


def parse_started_date(value):
    """'September 2 2026  8:02 AM' -> datetime.date(2026, 9, 2)"""
    value = " ".join(value.split())  # collapse double spaces
    dt = datetime.strptime(value, "%B %d %Y %I:%M %p")
    return dt.date()


def discover_quiz_csvs(script_dir):
    """Find all Moodle quiz grade reports dropped in the grades/ folder."""
    return sorted(glob.glob(os.path.join(script_dir, QUIZ_CSV_GLOB)))


def load_finishers(csv_path):
    """Return {(first_lower, last_lower)} of students with a Finished attempt,
    plus the single quiz date (as it appears in most rows)."""
    finishers = set()
    dates = Counter()
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            status = (row.get("Status") or "").strip()
            last = (row.get("Last name") or "").strip()
            first = (row.get("First name") or "").strip()
            started = (row.get("Started") or "").strip()
            if not last or not started:
                continue  # skips the trailing "Overall average" row
            if status != "Finished":
                continue
            try:
                dates[parse_started_date(started)] += 1
            except ValueError:
                continue
            finishers.add((first.lower(), last.lower()))
    if not dates:
        sys.exit(f"No 'Finished' attempts with a parseable Started date found in {csv_path}.")
    quiz_date = dates.most_common(1)[0][0]
    off_date_rows = sum(c for d, c in dates.items() if d != quiz_date)
    if off_date_rows:
        print(f"Note: {off_date_rows} finished attempt(s) fell on a different date than "
              f"the majority ({quiz_date.isoformat()}) and were still counted.")
    return finishers, quiz_date


def date_header_variants(d):
    """Excel date headers in this sheet look like '9/2/2026' (no zero-padding)."""
    return {f"{d.month}/{d.day}/{d.year}"}


def process_csv(ws, csv_path, dry_run, force):
    """Apply one quiz report's attendance to the already-open Attendance sheet.
    Returns True if any cell was (or would be) changed."""
    finishers, quiz_date = load_finishers(csv_path)
    print(f"Quiz date detected from CSV: {quiz_date.isoformat()}  "
          f"({len(finishers)} unique students finished)")

    header = [c.value for c in ws[1]]
    candidates = date_header_variants(quiz_date)
    col_idx = None
    for i, h in enumerate(header):
        if h in candidates:
            col_idx = i + 1  # openpyxl is 1-indexed
            break
    if col_idx is None:
        print(f"SKIPPED: no Attendance column found for date {quiz_date.isoformat()} "
              f"(looked for header in {sorted(candidates)}). "
              f"Existing headers: {header}")
        return False

    date_label = header[col_idx - 1]
    print(f"Writing to Attendance column '{date_label}' (column {col_idx})")

    matched_keys = set()
    changes = []  # (row, name, old, new)
    skipped = []  # (row, name, existing) -- already had a value, left alone
    for r in range(2, ws.max_row + 1):
        fn = ws.cell(r, 1).value
        ln = ws.cell(r, 2).value
        if fn is None and ln is None:
            continue
        key = ((fn or "").strip().lower(), (ln or "").strip().lower())
        present = key in finishers
        if present:
            matched_keys.add(key)
        new_val = "p" if present else "n"
        cell = ws.cell(r, col_idx)
        old_val = cell.value
        if old_val is not None and not force:
            if old_val != new_val:
                skipped.append((r, f"{fn} {ln}", old_val))
            continue
        if old_val != new_val:
            changes.append((r, f"{fn} {ln}", old_val, new_val))
        if not dry_run:
            cell.value = new_val

    unmatched_quiz = finishers - matched_keys
    print(f"{len(changes)} cell(s) {'would change' if dry_run else 'changed'}:")
    for r, name, old, new in changes:
        print(f"  row {r:>2}  {name:<28} {old!r:>6} -> {new!r}")

    if skipped:
        print(f"{len(skipped)} cell(s) already had a manually-entered value and were left "
              f"as-is (the quiz report disagreed — pass --force to overwrite if that's wrong):")
        for r, name, old in skipped:
            print(f"  row {r:>2}  {name:<28} kept {old!r}")

    if unmatched_quiz:
        print(f"WARNING: {len(unmatched_quiz)} student(s) finished the quiz per the CSV "
              f"but were not found on the Attendance roster (name mismatch or not enrolled "
              f"in this section) — nothing was written for them:")
        for first, last in sorted(unmatched_quiz, key=lambda x: x[1]):
            print(f"  {first.title()} {last.title()}")

    return bool(changes)


def main():
    argv = sys.argv[1:]
    dry_run = "--dry-run" in argv
    force = "--force" in argv
    argv = [a for a in argv if a not in ("--dry-run", "--force")]

    script_dir = os.path.dirname(os.path.abspath(__file__))
    xlsx_path = os.path.join(script_dir, "grades.xlsx")
    if not os.path.exists(xlsx_path):
        sys.exit(f"File not found: {xlsx_path}")

    if argv:
        csv_paths = argv
        for p in csv_paths:
            if not os.path.exists(p):
                sys.exit(f"File not found: {p}")
    else:
        csv_paths = discover_quiz_csvs(script_dir)
        if not csv_paths:
            sys.exit(f'No files matching "{QUIZ_CSV_GLOB}" found in {script_dir}.\n'
                      f"Drop the exported Moodle report(s) in this folder, or pass a CSV "
                      f"path explicitly.")
        print(f"No file given on the command line — found {len(csv_paths)} quiz report(s) "
              f"in {script_dir}:")
        for p in csv_paths:
            print(f"  {os.path.basename(p)}")
        print()

    wb = openpyxl.load_workbook(xlsx_path)
    if ATTENDANCE_SHEET not in wb.sheetnames:
        sys.exit(f"No '{ATTENDANCE_SHEET}' sheet in {xlsx_path}")
    ws = wb[ATTENDANCE_SHEET]

    any_changes = False
    for i, csv_path in enumerate(csv_paths):
        if len(csv_paths) > 1:
            print(f"=== {os.path.basename(csv_path)} ===")
        if process_csv(ws, csv_path, dry_run=dry_run, force=force):
            any_changes = True
        if i < len(csv_paths) - 1:
            print()

    if dry_run:
        print("\nDry run — no files were modified.")
        return

    if not any_changes:
        print("\nNo changes to save.")
        return

    backup_path = os.path.join(
        script_dir, f"grades_backup_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    )
    shutil.copy2(xlsx_path, backup_path)
    wb.save(xlsx_path)
    print(f"\nSaved {xlsx_path}")
    print(f"Backup of the previous version: {backup_path}")


if __name__ == "__main__":
    main()
